"""Document generation, authorization, and ownership tests."""

from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from sqlalchemy import func, select

from backend.auth import load_user_context
from backend.database import create_session_factory, create_sqlite_engine
from backend.models import AgentRun, GeneratedDocument
from backend.scripts.ingest_data import run_ingestion
from backend.services.documents import DocumentService
from backend.services.errors import PermissionDeniedError, ResourceNotFoundError


REPO_ROOT = Path(__file__).resolve().parents[1]


class DocumentServiceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = tempfile.TemporaryDirectory()
        root = Path(cls.temp_dir.name)
        cls.database = root / "documents.db"
        cls.storage = root / "storage"
        run_ingestion(REPO_ROOT / "data", cls.database)
        cls.engine = create_sqlite_engine(cls.database)
        cls.session_factory = create_session_factory(cls.engine)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.engine.dispose()
        cls.temp_dir.cleanup()

    def test_generates_real_xlsx_pdf_and_docx_energy_reports(self) -> None:
        generated = {}
        with self.session_factory.begin() as session:
            ctx = load_user_context(session, user_id="company_1_operator")
            service = DocumentService(session, self.storage)
            for file_format in ("xlsx", "pdf", "docx"):
                info = service.create_energy_report(ctx, file_format)
                generated[file_format] = service.download_document(
                    ctx, info.document_id, info.run_id
                )

        workbook = load_workbook(BytesIO(generated["xlsx"].content))
        self.assertEqual(workbook.sheetnames, ["Plants", "Energy Summary"])
        self.assertEqual(workbook["Plants"]["A1"].value, "Solar Plant Energy Report")
        self.assertTrue(generated["pdf"].content.startswith(b"%PDF"))
        word_document = Document(BytesIO(generated["docx"].content))
        self.assertEqual(
            word_document.paragraphs[0].text, "Solar Plant Energy Report"
        )
        for file_format, downloaded in generated.items():
            with self.subTest(file_format=file_format):
                self.assertGreater(downloaded.info.byte_size, 500)
                self.assertEqual(
                    downloaded.info.byte_size, len(downloaded.content)
                )
                self.assertEqual(downloaded.info.classification, "energy")

    def test_financial_user_can_generate_and_download_report(self) -> None:
        with self.session_factory.begin() as session:
            ctx = load_user_context(session, user_id="company_1_admin")
            service = DocumentService(session, self.storage)
            info = service.create_financial_report(ctx, "xlsx", month=3)
            downloaded = service.download_document(
                ctx, info.document_id, info.run_id
            )

        workbook = load_workbook(BytesIO(downloaded.content), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Market Price Summary", "Monthly Cost Summary"],
        )
        self.assertEqual(info.classification, "financial")

    def test_energy_only_user_cannot_generate_financial_report(self) -> None:
        with self.session_factory.begin() as session:
            ctx = load_user_context(session, user_id="company_2_operator")
            service = DocumentService(session, self.storage)
            before_documents = session.scalar(
                select(func.count()).select_from(GeneratedDocument)
            )
            before_runs = session.scalar(select(func.count()).select_from(AgentRun))
            with self.assertRaisesRegex(
                PermissionDeniedError, "financial data access is not permitted"
            ):
                service.create_financial_report(ctx, "pdf", month=3)
            after_documents = session.scalar(
                select(func.count()).select_from(GeneratedDocument)
            )
            after_runs = session.scalar(select(func.count()).select_from(AgentRun))

        self.assertEqual(after_documents, before_documents)
        self.assertEqual(after_runs, before_runs)

    def test_same_company_other_user_cannot_download_document(self) -> None:
        with self.session_factory.begin() as session:
            owner = load_user_context(session, user_id="company_1_operator")
            other_user = load_user_context(session, user_id="company_1_admin")
            service = DocumentService(session, self.storage)
            info = service.create_energy_report(owner, "pdf")

            with self.assertRaisesRegex(
                ResourceNotFoundError, "not found or not accessible"
            ):
                service.download_document(
                    other_user, info.document_id, info.run_id
                )

    def test_other_company_cannot_download_document(self) -> None:
        with self.session_factory.begin() as session:
            owner = load_user_context(session, user_id="company_1_admin")
            other_company = load_user_context(session, user_id="company_2_admin")
            service = DocumentService(session, self.storage)
            info = service.create_financial_report(owner, "docx", month=3)

            with self.assertRaisesRegex(
                ResourceNotFoundError, "not found or not accessible"
            ):
                service.download_document(
                    other_company, info.document_id, info.run_id
                )

    def test_wrong_run_id_cannot_download_document(self) -> None:
        with self.session_factory.begin() as session:
            owner = load_user_context(session, user_id="company_2_operator")
            service = DocumentService(session, self.storage)
            info = service.create_energy_report(owner, "xlsx")

            with self.assertRaisesRegex(
                ResourceNotFoundError, "not found or not accessible"
            ):
                service.download_document(
                    owner,
                    info.document_id,
                    "00000000-0000-0000-0000-000000000000",
                )

    def test_document_metadata_does_not_expose_storage_path(self) -> None:
        with self.session_factory.begin() as session:
            owner = load_user_context(session, user_id="company_2_operator")
            info = DocumentService(session, self.storage).create_energy_report(
                owner, "pdf"
            )

        self.assertFalse(hasattr(info, "storage_key"))
        self.assertFalse(hasattr(info, "company_id"))
        self.assertFalse(hasattr(info, "user_id"))


if __name__ == "__main__":
    unittest.main()
