"""Owned report generation and retrieval over scoped application services."""

from __future__ import annotations

import hashlib
import os
import tempfile
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, List, Optional, Sequence, Tuple

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.auth import UserContext
from backend.models import AgentRun, GeneratedDocument
from backend.services.errors import PermissionDeniedError, ResourceNotFoundError
from backend.services.finance import FinanceService
from backend.services.plants import PlantDataService


REPORT_FORMATS = {
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "pdf": ("application/pdf", "pdf"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
}


@dataclass(frozen=True)
class GeneratedDocumentInfo:
    document_id: str
    run_id: str
    filename: str
    media_type: str
    byte_size: int
    checksum: str
    classification: str
    created_at: datetime


@dataclass(frozen=True)
class DownloadedDocument:
    info: GeneratedDocumentInfo
    content: bytes


@dataclass(frozen=True)
class ReportTable:
    title: str
    headers: Tuple[str, ...]
    rows: Tuple[Tuple[Any, ...], ...]


@dataclass(frozen=True)
class ReportContent:
    title: str
    generated_at: datetime
    tables: Tuple[ReportTable, ...]


class DocumentService:
    def __init__(self, session: Session, storage_root: Path) -> None:
        self._session = session
        self._storage_root = storage_root.expanduser().resolve()

    def create_energy_report(
        self,
        ctx: UserContext,
        file_format: str,
        start: Optional[datetime] = None,
        end: Optional[datetime] = None,
    ) -> GeneratedDocumentInfo:
        """Create an owned operational report for all accessible plants."""
        plant_service = PlantDataService(self._session)
        plants = plant_service.list_plants(ctx)
        summaries = plant_service.get_energy_summary(ctx, start=start, end=end)
        plant_rows = tuple(
            (
                plant.plant_id,
                plant.name,
                _display_decimal(plant.nominal_power_kw),
                plant.region or "",
                plant.commissioning_date.isoformat()
                if plant.commissioning_date
                else "",
            )
            for plant in plants
        )
        metric_rows = tuple(
            (
                summary.plant_id,
                summary.plant_name,
                metric.datasource_id,
                metric.name,
                metric.unit,
                metric.aggregation,
                metric.sample_count,
                _display_decimal(metric.value),
            )
            for summary in summaries
            for metric in summary.metrics
        )
        content = ReportContent(
            title="Solar Plant Energy Report",
            generated_at=datetime.now(timezone.utc),
            tables=(
                ReportTable(
                    title="Plants",
                    headers=(
                        "Plant ID",
                        "Plant",
                        "Nominal Power (kW)",
                        "Region",
                        "Commissioning Date",
                    ),
                    rows=plant_rows,
                ),
                ReportTable(
                    title="Energy Summary",
                    headers=(
                        "Plant ID",
                        "Plant",
                        "Datasource ID",
                        "Metric",
                        "Unit",
                        "Aggregation",
                        "Samples",
                        "Value",
                    ),
                    rows=metric_rows,
                ),
            ),
        )
        return self._create_document(ctx, file_format, "energy", content)

    def create_financial_report(
        self,
        ctx: UserContext,
        file_format: str,
        month: Optional[int] = None,
    ) -> GeneratedDocumentInfo:
        """Create an owned financial report after service authorization."""
        finance = FinanceService(self._session)
        # Authorization happens inside FinanceService before either query.
        market_prices = finance.get_market_price_summary(ctx)
        monthly_costs = finance.get_monthly_cost_summary(ctx, month=month)
        price_rows = tuple(
            (
                item.zone,
                item.sample_count,
                _scaled_value(item.minimum_eur_per_mwh_micros, 1_000_000),
                _scaled_value(item.maximum_eur_per_mwh_micros, 1_000_000),
                _scaled_value(item.average_eur_per_mwh_micros, 1_000_000),
            )
            for item in market_prices
        )
        cost_rows = tuple(
            (
                item.year,
                item.month,
                item.category,
                _scaled_value(item.amount_eur_minor, 100),
            )
            for item in monthly_costs
        )
        content = ReportContent(
            title="Solar Plant Financial Report",
            generated_at=datetime.now(timezone.utc),
            tables=(
                ReportTable(
                    title="Market Price Summary",
                    headers=(
                        "Zone",
                        "Samples",
                        "Minimum EUR/MWh",
                        "Maximum EUR/MWh",
                        "Average EUR/MWh",
                    ),
                    rows=price_rows,
                ),
                ReportTable(
                    title="Monthly Cost Summary",
                    headers=("Year", "Month", "Category", "Amount EUR"),
                    rows=cost_rows,
                ),
            ),
        )
        return self._create_document(ctx, file_format, "financial", content)

    def download_document(
        self, ctx: UserContext, document_id: str, run_id: str
    ) -> DownloadedDocument:
        """Return bytes only when all ownership dimensions match."""
        document = self._session.scalar(
            select(GeneratedDocument).where(
                GeneratedDocument.document_id == document_id,
                GeneratedDocument.company_id == ctx.company_id,
                GeneratedDocument.user_id == ctx.user_id,
                GeneratedDocument.run_id == run_id,
            )
        )
        if document is None:
            raise ResourceNotFoundError("document not found or not accessible")
        if document.classification == "financial" and not ctx.can_view_financials:
            raise PermissionDeniedError("financial data access is not permitted")

        path = self._resolve_storage_key(document.storage_key)
        try:
            content = path.read_bytes()
        except OSError as error:
            raise ResourceNotFoundError(
                "document not found or not accessible"
            ) from error
        if (
            len(content) != document.byte_size
            or hashlib.sha256(content).hexdigest() != document.checksum
        ):
            raise ResourceNotFoundError("document not found or not accessible")
        return DownloadedDocument(info=_document_info(document), content=content)

    def _create_document(
        self,
        ctx: UserContext,
        file_format: str,
        classification: str,
        content: ReportContent,
    ) -> GeneratedDocumentInfo:
        normalized_format = file_format.strip().lower()
        if normalized_format not in REPORT_FORMATS:
            raise ValueError(
                f"unsupported report format; expected one of {sorted(REPORT_FORMATS)}"
            )

        media_type, extension = REPORT_FORMATS[normalized_format]
        generated = _render_report(content, normalized_format)
        checksum = hashlib.sha256(generated).hexdigest()
        run_id = str(uuid.uuid4())
        document_id = str(uuid.uuid4())
        filename = f"{classification}-report-{document_id[:8]}.{extension}"
        storage_key = self._storage_key(ctx, run_id, document_id, extension)
        destination = self._resolve_storage_key(storage_key)
        now = datetime.now(timezone.utc)

        destination.parent.mkdir(parents=True, exist_ok=True)
        temp_path: Optional[Path] = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=destination.parent, prefix=".pending-", delete=False
            ) as handle:
                handle.write(generated)
                temp_path = Path(handle.name)
            os.replace(temp_path, destination)
            temp_path = None

            run = AgentRun(
                run_id=run_id,
                company_id=ctx.company_id,
                user_id=ctx.user_id,
                run_type=f"{classification}_report",
                status="completed",
                result_text=f"Generated {filename}",
                created_at=now,
                completed_at=now,
            )
            self._session.add(run)
            self._session.flush()
            document = GeneratedDocument(
                document_id=document_id,
                company_id=ctx.company_id,
                user_id=ctx.user_id,
                run_id=run_id,
                storage_key=storage_key,
                media_type=media_type,
                filename=filename,
                byte_size=len(generated),
                checksum=checksum,
                classification=classification,
                created_at=now,
            )
            self._session.add(document)
            self._session.flush()
            return _document_info(document)
        except Exception:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
            destination.unlink(missing_ok=True)
            raise

    def _storage_key(
        self,
        ctx: UserContext,
        run_id: str,
        document_id: str,
        extension: str,
    ) -> str:
        company_segment = hashlib.sha256(ctx.company_id.encode()).hexdigest()[:20]
        user_segment = hashlib.sha256(ctx.user_id.encode()).hexdigest()[:20]
        return (
            f"{company_segment}/{user_segment}/{run_id}/"
            f"{document_id}.{extension}"
        )

    def _resolve_storage_key(self, storage_key: str) -> Path:
        path = (self._storage_root / storage_key).resolve()
        if self._storage_root not in path.parents:
            raise ResourceNotFoundError("document not found or not accessible")
        return path


def _render_report(content: ReportContent, file_format: str) -> bytes:
    if file_format == "xlsx":
        return _render_xlsx(content)
    if file_format == "docx":
        return _render_docx(content)
    if file_format == "pdf":
        return _render_pdf(content)
    raise ValueError("unsupported report format")


def _render_xlsx(content: ReportContent) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for report_table in content.tables:
        worksheet = workbook.create_sheet(report_table.title[:31])
        worksheet.append([content.title])
        worksheet.append(["Generated at", content.generated_at.isoformat()])
        worksheet.append([])
        worksheet.append(list(report_table.headers))
        for cell in worksheet[4]:
            cell.font = Font(bold=True)
        for row in report_table.rows:
            worksheet.append([_spreadsheet_value(value) for value in row])
        for column in range(1, len(report_table.headers) + 1):
            values = [
                str(worksheet.cell(row=row, column=column).value or "")
                for row in range(1, worksheet.max_row + 1)
            ]
            worksheet.column_dimensions[get_column_letter(column)].width = min(
                max(len(value) for value in values) + 2, 40
            )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _render_docx(content: ReportContent) -> bytes:
    document = Document()
    document.add_heading(content.title, level=0)
    document.add_paragraph(f"Generated at: {content.generated_at.isoformat()}")
    for report_table in content.tables:
        document.add_heading(report_table.title, level=1)
        table = document.add_table(rows=1, cols=len(report_table.headers))
        table.style = "Table Grid"
        for index, header in enumerate(report_table.headers):
            table.rows[0].cells[index].text = header
        for row in report_table.rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_pdf(content: ReportContent) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story: List[Any] = [
        Paragraph(content.title, styles["Title"]),
        Paragraph(
            f"Generated at: {content.generated_at.isoformat()}",
            styles["BodyText"],
        ),
        Spacer(1, 12),
    ]
    for report_table in content.tables:
        story.append(Paragraph(report_table.title, styles["Heading2"]))
        data = [list(report_table.headers)] + [
            [str(value) for value in row] for row in report_table.rows
        ]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.extend([table, Spacer(1, 12)])
    document.build(story)
    return buffer.getvalue()


def _document_info(document: GeneratedDocument) -> GeneratedDocumentInfo:
    return GeneratedDocumentInfo(
        document_id=document.document_id,
        run_id=document.run_id,
        filename=document.filename,
        media_type=document.media_type,
        byte_size=document.byte_size,
        checksum=document.checksum,
        classification=document.classification,
        created_at=document.created_at,
    )


def _display_decimal(value: Optional[Decimal]) -> str:
    return "" if value is None else format(value, "f")


def _scaled_value(value: int, scale: int) -> str:
    return format(Decimal(value) / Decimal(scale), "f")


def _spreadsheet_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value
